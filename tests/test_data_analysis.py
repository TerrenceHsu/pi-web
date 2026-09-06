"""Small real-data contracts for the optional analysis product (no LLM/network)."""

from __future__ import annotations

import asyncio
import io
import json
import time
from pathlib import Path

import pytest
from fastapi import UploadFile
from fastapi.testclient import TestClient

from agent_workspace import WorkspaceStore
from agent_workspace.analysis import AnalysisError, AnalysisRequest
from coding_agent_app.data_analysis.engine import calculate
from coding_agent_app.data_analysis.service import DataAnalysisService
from pi_agent_core_py.agent import Agent
from pi_agent_core_py.harness import AgentHarness
from pi_agent_core_py.messages import ToolCall
from pi_agent_core_py.model_client import DoneEvent, FakeClient, TextDeltaEvent, ToolCallEvent
from pi_agent_core_py.telemetry import InMemoryTelemetryContext
from pi_agent_core_py.web.app import create_app, dispose_app

pytest.importorskip("pandas")
pytest.importorskip("matplotlib")


def _data(tmp_path: Path) -> Path:
    path = tmp_path / "source.csv"
    path.write_text(
        "id,region,sales,date\n001,East,10,2026-01-01\n"
        "002,West,20,2026-01-02\n003,East,30,2026-02-01\n"
        "004,East,,2026-02-02\n",
        encoding="utf-8",
    )
    return path


@pytest.mark.parametrize(
    "action,options,expected",
    [
        ("inspect", {}, ["001", "East", "10", "2026-01-01"]),
        (
            "aggregate",
            {"group_by": ["region"], "metrics": [{"column": "sales", "operation": "sum"}]},
            ["East", 40.0],
        ),
        (
            "timeseries",
            {"date_column": "date", "metrics": [{"column": "sales", "operation": "sum"}]},
            ["2026-01", 30.0],
        ),
    ],
)
def test_fixed_operations_use_full_input_and_preserve_identifiers(
    tmp_path, action, options, expected
):
    request = AnalysisRequest(file_id="file-test", action=action, **options)
    result = calculate(_data(tmp_path), request, tmp_path)
    assert result["source_rows"] == result["analyzed_rows"] == 4
    assert result["rows"][0] == expected
    assert result["schema"][0] == {"name": "id", "type": "text"}


def test_profile_filter_top_n_and_invalid_numeric_data(tmp_path):
    source = _data(tmp_path)
    result = calculate(source, AnalysisRequest(file_id="file-a", action="profile"), tmp_path)
    sales = dict(zip(result["columns"], result["rows"][2], strict=True))
    assert sales["missing"] == 1 and sales["mean"] == 20
    result = calculate(
        source,
        AnalysisRequest.model_validate(
            {
                "file_id": "file-a",
                "action": "aggregate",
                "group_by": ["region"],
                "filters": [{"column": "sales", "operation": "gte", "value": 20}],
                "metrics": [{"column": "sales", "operation": "sum"}],
                "sort_by": "sales__sum",
                "limit": 1,
            }
        ),
        tmp_path,
    )
    assert result["rows"] == [["East", 30]] and result["limited"]
    assert result["analyzed_rows"] == 2
    with pytest.raises(AnalysisError, match="invalid_numeric_values"):
        calculate(
            source,
            AnalysisRequest.model_validate(
                {
                    "file_id": "file-a",
                    "action": "aggregate",
                    "metrics": [{"column": "region", "operation": "sum"}],
                }
            ),
            tmp_path,
        )


def test_inspect_numeric_sort_keeps_original_cells_and_identifiers(tmp_path):
    source = tmp_path / "source.csv"
    source.write_text("id,value\n001,2\n002,10\n003,100\n", encoding="utf-8")
    result = calculate(
        source,
        AnalysisRequest(
            file_id="file-a",
            action="inspect",
            sort_by="value",
            limit=2,
        ),
        tmp_path,
    )
    assert result["rows"] == [["003", "100"], ["002", "10"]]
    assert result["limited"] and result["source_rows"] == 3


def test_xlsx_sheet_selection_and_missing_formula_cache_fail_closed(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Data"
    book.active.append(["id", "value"])
    book.active.append(["001", 3])
    second = book.create_sheet("Formula")
    second.append(["total"])
    second.append(["=1+2"])
    path = tmp_path / "source.xlsx"
    book.save(path)
    result = calculate(
        path, AnalysisRequest(file_id="file-a", action="inspect", sheet="Data"), tmp_path
    )
    assert result["rows"] == [["001", 3]]
    assert result["source"]["sheets"] == ["Data", "Formula"]
    with pytest.raises(AnalysisError, match="formula_cache_missing"):
        calculate(
            path,
            AnalysisRequest.model_validate(
                {
                    "file_id": "file-a",
                    "action": "aggregate",
                    "sheet": "Formula",
                    "metrics": [{"column": "total", "operation": "sum"}],
                }
            ),
            tmp_path,
        )


@pytest.mark.parametrize("kind", ["bar", "line", "histogram", "scatter"])
def test_real_chart_is_png_and_parquet_is_supported(tmp_path, kind):
    import pyarrow as pa
    import pyarrow.parquet as pq

    path = tmp_path / "source.parquet"
    pq.write_table(pa.table({"x": [1, 2, 3], "y": [4, 5, 6]}), path)
    result = calculate(
        path,
        AnalysisRequest(
            file_id="file-a",
            action="chart",
            chart_type=kind,
            x="x",
            y="y",
        ),
        tmp_path,
    )
    assert result["has_chart"] and result["analyzed_rows"] == 3
    assert (tmp_path / "chart.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")


def _wait(client: TestClient, url: str) -> dict:
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        response = client.get(url)
        assert response.status_code == 200, response.text
        payload = response.json()
        if payload["status"] not in {"queued", "running"}:
            return payload
        time.sleep(0.05)
    pytest.fail("analysis did not finish")


def test_optional_tool_actual_worker_save_reload_and_session_isolation(tmp_path):
    def app_factory():
        return create_app(
            AgentHarness(Agent(system_prompt="", client=FakeClient([]))),
            db_path=tmp_path / "session.sqlite",
            uploads_dir=tmp_path / "uploads",
        )

    app = app_factory()
    with TestClient(app) as client:
        sid = client.post("/api/sessions", json={"title": "analysis"}).json()["id"]
        other = client.post("/api/sessions", json={"title": "other"}).json()["id"]
        base = f"/api/workspaces/{sid}"
        uploaded = client.post(
            f"/api/sessions/{sid}/files",
            files={
                "files": ("sales.csv", _data(tmp_path).read_bytes(), "text/csv"),
            },
        ).json()["files"][0]
        args = {
            "file_id": uploaded["id"],
            "action": "chart",
            "chart_type": "bar",
            "x": "region",
            "y": "sales",
        }
        assert client.post(base + "/analysis", json=args).status_code == 409
        selected = client.put(
            base + "/extensions",
            json={
                "mcp_server_names": [],
                "skill_names": [],
                "tool_names": ["analyze_data"],
            },
        )
        assert selected.status_code == 200, selected.text
        assert selected.json()["selected_tool_names"] == ["analyze_data"]
        started = client.post(base + "/analysis", json=args)
        assert started.status_code == 202, started.text
        run_id = started.json()["id"]
        url = base + f"/analysis/{run_id}"
        result = _wait(client, url)
        assert result["status"] == "succeeded", result
        assert client.get(url + "/chart").content.startswith(b"\x89PNG")
        assert client.get(f"/api/workspaces/{other}/analysis/{run_id}").status_code == 404
        before = client.get(f"/api/sessions/{sid}/workspace").json()["workspace"]["revision"]
        saved = client.post(url + "/save")
        assert saved.status_code == 200, saved.text
        assert len(saved.json()["files"]) == 4
        assert all("path" not in ref for ref in saved.json()["files"])
        after = client.get(f"/api/sessions/{sid}/workspace").json()["workspace"]["revision"]
        assert after == before + 1
        assert client.post(url + "/save").status_code == 200
        assert (
            client.get(f"/api/sessions/{sid}/files/{uploaded['id']}").content
            == _data(tmp_path).read_bytes()
        )
        assert client.put(
            base + "/extensions",
            json={
                "mcp_server_names": [],
                "skill_names": [],
            },
        ).json()["selected_tool_names"] == ["analyze_data"]
        assert (
            client.put(
                base + "/extensions",
                json={
                    "mcp_server_names": [],
                    "skill_names": [],
                    "tool_names": [],
                },
            ).status_code
            == 200
        )
        assert client.post(base + "/analysis", json=args).status_code == 409
    dispose_app(app)
    reopened = app_factory()
    with TestClient(reopened) as client:
        assert client.get(url).json()["result"]["source_sha256"] == uploaded["sha256"]
        assert client.get(base + "/extensions").json()["selected_tool_names"] == []
        assert client.get(url + "/chart").status_code == 200
        assert client.delete(f"/api/sessions/{sid}").status_code == 200
        assert client.get(url).status_code == 404
    dispose_app(reopened)


def test_web_agent_uses_selected_tool_and_request_scoped_workspace(tmp_path):
    call = ToolCall(
        id="analysis-call",
        name="analyze_data",
        arguments={
            "file_id": "pending",
            "action": "inspect",
        },
    )
    fake = FakeClient(
        [
            [ToolCallEvent(tool_call=call), DoneEvent(stop_reason="tool_use")],
            [TextDeltaEvent(delta="Analysis ready."), DoneEvent(stop_reason="stop")],
        ]
    )
    app = create_app(
        AgentHarness(Agent(system_prompt="", client=fake)),
        db_path=tmp_path / "session.sqlite",
        uploads_dir=tmp_path / "uploads",
    )
    with TestClient(app) as client:
        sid = client.post("/api/sessions", json={"title": "Agent analysis"}).json()["id"]
        uploaded = client.post(
            f"/api/sessions/{sid}/files",
            files={
                "files": ("table.csv", b"x,y\n001,3\n", "text/csv"),
            },
        ).json()["files"][0]
        call.arguments["file_id"] = uploaded["id"]
        base = f"/api/workspaces/{sid}"
        assert client.put(base + "/extensions", json={"tool_names": ["analyze_data"]}).is_success
        response = client.post(
            "/api/prompt",
            json={
                "session_id": sid,
                "text": "Please inspect the uploaded table using analyze_data.",
            },
        )
        assert response.status_code == 200, response.text
        runs = client.get(base + "/analysis").json()["runs"]
        assert len(runs) == 1 and runs[0]["status"] == "succeeded", response.text
        result = client.get(base + "/analysis/" + runs[0]["id"]).json()["result"]
        assert result["rows"] == [["001", "3"]]
        assert result["session_id"] == sid
        assert any(tool.name == "analyze_data" for tool in fake.last_tools)
        assert not any(
            file["logical_path"].startswith("artifacts/analysis/")
            for file in client.get(f"/api/sessions/{sid}/files").json()["files"]
        )
    dispose_app(app)


@pytest.mark.asyncio
async def test_worker_timeout_and_crash_recovery(tmp_path):
    store = WorkspaceStore(tmp_path / "files")
    source = await store.save(
        "one", UploadFile(filename="data.csv", file=io.BytesIO(b"x,y\n1,2\n"))
    )

    async def enabled(_sid):
        return True

    service = DataAnalysisService(tmp_path / "analysis", store, enabled, timeout_seconds=0.01)
    await service.init()
    request = AnalysisRequest(file_id=source.id, action="inspect")
    with pytest.raises(AnalysisError, match="analysis_timeout"):
        await service.analyze("one", request)
    run_id = (await service.list_runs("one"))[0]["id"]
    await service._update(run_id, "running")
    await service.close()
    reopened = DataAnalysisService(tmp_path / "analysis", store, enabled)
    await reopened.init()
    try:
        assert (await reopened.get("one", run_id))["status"] == "interrupted"
        new_id = await reopened.start("one", request)
        assert (await reopened.cancel("one", new_id))["status"] == "cancelled"
    finally:
        await reopened.close()


def test_tsv_formula_injection_and_formula_cache_profile_guard(tmp_path):
    path = tmp_path / "source.tsv"
    path.write_text("name\tvalue\n=1+1\t2\n@SUM(A1)\t3\n", encoding="utf-8")
    result = calculate(path, AnalysisRequest(file_id="file-a", action="inspect"), tmp_path)
    assert result["rows"][0] == ["=1+1", "2"]
    assert "'=1+1" in (tmp_path / "result.csv").read_text(encoding="utf-8-sig")
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["x", "total"])
    book.active.append([1, "=1+2"])
    path = tmp_path / "source.xlsx"
    book.save(path)
    with pytest.raises(AnalysisError, match="formula_cache_missing"):
        calculate(
            path,
            AnalysisRequest.model_validate(
                {
                    "file_id": "file-a",
                    "action": "profile",
                    "filters": [{"column": "x", "operation": "gte", "value": 1}],
                }
            ),
            tmp_path,
        )


async def test_parallel_sessions_idempotent_save_tamper_and_metadata_only_telemetry(tmp_path):
    store = WorkspaceStore(tmp_path / "files")
    for session_id in ("one", "two"):
        await store.ensure_session_workspace(session_id)
    refs = [
        await store.save(
            sid,
            UploadFile(
                filename="private.csv",
                file=io.BytesIO(f"private_column,value\nsecret,{n}\n".encode()),
            ),
        )
        for sid, n in (("one", 2), ("two", 9))
    ]

    async def enabled(_sid):
        return True

    service = DataAnalysisService(tmp_path / "analysis", store, enabled)
    service.telemetry = recorder = InMemoryTelemetryContext()
    await service.init()
    try:
        results = await asyncio.gather(
            *[
                service.analyze(ref.session_id, AnalysisRequest(file_id=ref.id, action="inspect"))
                for ref in refs
            ]
        )
        assert [r["rows"][0][1] for r in results] == ["2", "9"]
        rid = results[0]["run_id"]
        before_files = len(await store.list_session("one"))
        saved = await asyncio.gather(service.save("one", rid), service.save("one", rid))
        assert saved[0] == saved[1]
        assert len(await store.list_session("one")) == before_files + 3
        manifest_path = await service.output("one", rid, "manifest.json")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assert manifest["engine_versions"]["pandas"]
        csv_path = await service.output("one", rid, "result.csv")
        csv_path.write_bytes(b"tampered")
        with pytest.raises(AnalysisError, match="analysis_output_changed"):
            await service.save("one", rid)
        with pytest.raises(AnalysisError, match="source_not_found"):
            await service.start("two", AnalysisRequest(file_id=refs[0].id, action="inspect"))
        encoded = repr(recorder.get_spans())
        assert "data_analysis.run" in encoded and "rows_processed" in encoded
        assert "private_column" not in encoded and "secret" not in encoded
        await service.delete_session("one")
        # Explicit exports survive history deletion.
        assert len(await store.list_session("one")) == before_files + 3
        with pytest.raises(AnalysisError, match="analysis_unavailable"):
            await service.start("one", AnalysisRequest(file_id=refs[0].id, action="inspect"))
    finally:
        await service.close()
