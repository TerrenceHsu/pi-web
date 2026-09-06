"""Fixed Workspace document conversion, isolated from Knowledge/LLM Wiki.

The source file remains the only fact source.  Converters read an immutable
Workspace snapshot and return deterministic files; :class:`WorkspaceStore`
publishes the complete generated set in one revision-bound transaction.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import re
import shutil
import time
import uuid
import zipfile
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path, PurePosixPath
from typing import Any, Literal, Protocol, cast
from urllib.parse import urlsplit

from pydantic import BaseModel, ConfigDict, Field, model_validator

from .store import (
    FileRef,
    FileStoreError,
    WorkspacePublishChange,
    WorkspacePublishPolicyError,
    WorkspaceStore,
    is_document_conversion_workspace_path,
    normalize_workspace_logical_path,
    workspace_document_output_root,
)

WORKSPACE_DOCUMENT_MANIFEST_SCHEMA: Literal["pi-agent-workspace-document/v1"] = (
    "pi-agent-workspace-document/v1"
)
WORKSPACE_DOCUMENT_CONFIG_VERSION = "workspace-documents-2026-08-29"
_MAX_ZIP_ENTRIES = 10_000
_MAX_ZIP_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
_MAX_GENERATED_FILES = 256
_MAX_GENERATED_BYTES = 100 * 1024 * 1024
_MAX_PDF_PAGES = 5_000
_MAX_PDF_TEXT_CHARS = 25_000_000
_MAX_WORKBOOK_SHEETS = 100
_MAX_WORKBOOK_CELLS = 1_000_000

DocumentConversionStatus = Literal["succeeded", "needs_ocr", "failed"]


class WorkspaceDocumentError(FileStoreError):
    """Base error for the fixed Workspace document pipeline."""


class UnsupportedWorkspaceDocumentError(WorkspaceDocumentError):
    """The source is not a supported immutable Workspace document."""


class WorkspaceDocumentDependencyError(WorkspaceDocumentError):
    """A fixed converter dependency is not installed."""


class InvalidWorkspaceDocumentError(WorkspaceDocumentError):
    """The source is malformed, encrypted, or exceeds safe parse limits."""


class GeneratedDocumentFile(BaseModel):
    """One converter output relative to a document root."""

    model_config = ConfigDict(frozen=True, extra="forbid", arbitrary_types_allowed=True)

    relative_path: str
    content: bytes
    mime: str

    @model_validator(mode="after")
    def _validate_relative_path(self) -> GeneratedDocumentFile:
        path = PurePosixPath(self.relative_path)
        if (
            path.is_absolute()
            or not path.parts
            or any(part in {"", ".", ".."} for part in path.parts)
            or "\\" in self.relative_path
        ):
            raise ValueError("generated document path must be a safe relative path")
        return self


class DocumentConversionPayload(BaseModel):
    """Provider-neutral fixed converter result before persistence."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    status: Literal["succeeded", "needs_ocr"]
    files: tuple[GeneratedDocumentFile, ...]
    warnings: tuple[str, ...] = ()


class WorkspaceDocumentArtifact(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")

    logical_path: str
    mime: str
    size: int = Field(ge=0)
    sha256: str = Field(pattern=r"^[0-9a-f]{64}$")


class WorkspaceDocumentManifest(BaseModel):
    """Canonical, path-safe evidence for one fixed conversion."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    schema_version: Literal["pi-agent-workspace-document/v1"] = WORKSPACE_DOCUMENT_MANIFEST_SCHEMA
    document_id: str
    source_file_id: str
    source_name: str
    source_logical_path: str
    source_mime: str
    source_size: int = Field(ge=0)
    source_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    converter_name: str
    converter_version: str
    config_version: str = WORKSPACE_DOCUMENT_CONFIG_VERSION
    status: DocumentConversionStatus
    converted_at_ms: int = Field(ge=0)
    warnings: tuple[str, ...] = ()
    error_code: str | None = None
    generated_files: tuple[WorkspaceDocumentArtifact, ...] = ()


class WorkspaceDocumentConversionResult(BaseModel):
    """API-safe conversion outcome and the files now visible in Workspace."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    source_file_id: str
    document_id: str
    status: DocumentConversionStatus
    reused: bool = False
    workspace_revision: int = Field(ge=0)
    manifest_file_id: str | None = None
    primary_file_id: str | None = None
    files: tuple[FileRef, ...] = ()
    warnings: tuple[str, ...] = ()
    error_code: str | None = None


class WorkspaceDocumentConverter(Protocol):
    """No-I/O-policy interface implemented by fixed local converters."""

    @property
    def source_extension(self) -> str: ...

    @property
    def converter_name(self) -> str: ...

    @property
    def converter_version(self) -> str: ...

    def convert(self, source_path: Path, *, source_name: str) -> DocumentConversionPayload: ...


def _package_version(distribution: str) -> str:
    try:
        return version(distribution)
    except PackageNotFoundError:
        return "unavailable"


def _markdown_title(source_name: str) -> str:
    return PurePosixPath(source_name).stem.replace("#", "").strip() or "Document"


def _zip_preflight(source_path: Path) -> None:
    try:
        with zipfile.ZipFile(source_path) as archive:
            members = archive.infolist()
            if len(members) > _MAX_ZIP_ENTRIES:
                raise InvalidWorkspaceDocumentError("document archive has too many members")
            expanded = 0
            for member in members:
                if member.flag_bits & 0x1:
                    raise InvalidWorkspaceDocumentError(
                        "encrypted document archives are unsupported"
                    )
                expanded += member.file_size
                if expanded > _MAX_ZIP_UNCOMPRESSED_BYTES:
                    raise InvalidWorkspaceDocumentError(
                        "document archive exceeds the uncompressed size limit"
                    )
    except zipfile.BadZipFile as exc:
        raise InvalidWorkspaceDocumentError("document is not a valid ZIP package") from exc


def _safe_asset_name(raw_name: str, fallback: str) -> str:
    name = PurePosixPath(raw_name.replace("\\", "/")).name
    name = re.sub(r"[^\w.\-]+", "-", name, flags=re.UNICODE).strip(".-")
    return (name or fallback)[:180]


class PypdfWorkspaceConverter:
    source_extension = ".pdf"
    converter_name = "workspace-pypdf"

    @property
    def converter_version(self) -> str:
        return _package_version("pypdf")

    def convert(self, source_path: Path, *, source_name: str) -> DocumentConversionPayload:
        try:
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
        except ImportError as exc:  # pragma: no cover - dependency gate
            raise WorkspaceDocumentDependencyError(
                "pypdf is required for Workspace PDF conversion"
            ) from exc
        try:
            reader = PdfReader(source_path)
        except (PdfReadError, OSError, ValueError) as exc:
            raise InvalidWorkspaceDocumentError("PDF could not be parsed") from exc
        if reader.is_encrypted:
            raise InvalidWorkspaceDocumentError("encrypted PDFs are unsupported")
        if not reader.pages:
            raise InvalidWorkspaceDocumentError("PDF contains no pages")
        if len(reader.pages) > _MAX_PDF_PAGES:
            raise InvalidWorkspaceDocumentError("PDF has too many pages")

        sections = [f"# {_markdown_title(source_name)}", ""]
        warnings: list[str] = []
        assets: list[GeneratedDocumentFile] = []
        non_whitespace = 0
        text_chars = 0
        for page_number, page in enumerate(reader.pages, start=1):
            sections.extend((f"## Page {page_number}", ""))
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except Exception:
                text = ""
                warnings.append(f"page_{page_number}_text_extraction_failed")
            text = text.replace("\x00", "").strip()
            text_chars += len(text)
            if text_chars > _MAX_PDF_TEXT_CHARS:
                raise InvalidWorkspaceDocumentError("PDF extracted text exceeds the limit")
            non_whitespace += sum(not char.isspace() for char in text)
            sections.extend((text or "_No extractable text on this page._", ""))
            try:
                images: Iterable[object] = page.images
                for image_number, image in enumerate(images, start=1):
                    if len(assets) >= _MAX_GENERATED_FILES - 2:
                        warnings.append("pdf_image_limit_reached")
                        break
                    image_value = cast(Any, image)
                    data = bytes(image_value.data)
                    raw_name = str(getattr(image, "name", ""))
                    filename = _safe_asset_name(
                        raw_name,
                        f"page-{page_number:04d}-image-{image_number:03d}.bin",
                    )
                    if "." not in filename:
                        filename += ".bin"
                    assets.append(
                        GeneratedDocumentFile(
                            relative_path=f"assets/page-{page_number:04d}-{filename}",
                            content=data,
                            mime="application/octet-stream",
                        )
                    )
            except Exception:
                warnings.append(f"page_{page_number}_image_extraction_failed")

        status: Literal["succeeded", "needs_ocr"] = (
            "succeeded" if non_whitespace >= max(20, len(reader.pages) * 5) else "needs_ocr"
        )
        if status == "needs_ocr":
            warnings.append("needs_ocr")
            sections.insert(
                2, "> OCR is not enabled; this PDF has insufficient extractable text.\n"
            )
        warnings.append("pdf_tables_not_structurally_extracted")
        content = "\n".join(sections).rstrip() + "\n"
        return DocumentConversionPayload(
            status=status,
            files=(
                GeneratedDocumentFile(
                    relative_path="content.md",
                    content=content.encode("utf-8"),
                    mime="text/markdown",
                ),
                *assets,
            ),
            warnings=tuple(dict.fromkeys(warnings)),
        )


def _escape_markdown_cell(value: object) -> str:
    return str(value if value is not None else "").replace("\n", "<br>").replace("|", "\\|")


def _safe_markdown_link(label: str, url: str) -> str:
    normalized = url.strip()
    try:
        parsed = urlsplit(normalized)
    except ValueError:
        return label
    if not (normalized.startswith("#") or parsed.scheme.casefold() in {"http", "https", "mailto"}):
        return label
    safe_label = label.replace("[", "\\[").replace("]", "\\]")
    safe_url = normalized.replace("(", "%28").replace(")", "%29")
    return f"[{safe_label}]({safe_url})"


class DocxWorkspaceConverter:
    source_extension = ".docx"
    converter_name = "workspace-python-docx"

    @property
    def converter_version(self) -> str:
        return _package_version("python-docx")

    def convert(self, source_path: Path, *, source_name: str) -> DocumentConversionPayload:
        _zip_preflight(source_path)
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:  # pragma: no cover - dependency gate
            raise WorkspaceDocumentDependencyError(
                "python-docx is required for Workspace DOCX conversion"
            ) from exc
        try:
            document = Document(str(source_path))
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            raise InvalidWorkspaceDocumentError("DOCX could not be parsed") from exc

        lines = [f"# {_markdown_title(source_name)}", ""]
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                fragments: list[str] = []
                for item in paragraph.iter_inner_content():
                    item_text = str(item.text)
                    item_url = str(getattr(item, "url", "") or "")
                    fragments.append(
                        _safe_markdown_link(item_text, item_url) if item_url else item_text
                    )
                text = "".join(fragments).strip()
                if not text:
                    continue
                style_name = str(getattr(paragraph.style, "name", "") or "")
                heading = re.fullmatch(r"Heading\s+([1-6])", style_name, re.IGNORECASE)
                if heading:
                    lines.extend((f"{'#' * int(heading.group(1))} {text}", ""))
                elif "list" in style_name.casefold():
                    lines.append(f"- {text}")
                else:
                    lines.extend((text, ""))
            elif child.tag.endswith("}tbl"):
                table = Table(child, document)
                rows = [
                    [_escape_markdown_cell(cell.text.strip()) for cell in row.cells]
                    for row in table.rows
                ]
                if not rows:
                    continue
                width = max(len(row) for row in rows)
                normalized = [row + [""] * (width - len(row)) for row in rows]
                lines.append("| " + " | ".join(normalized[0]) + " |")
                lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
                for row in normalized[1:]:
                    lines.append("| " + " | ".join(row) + " |")
                lines.append("")

        assets: list[GeneratedDocumentFile] = []
        seen_parts: set[str] = set()
        for relationship in document.part.rels.values():
            if bool(getattr(relationship, "is_external", False)):
                continue
            target = getattr(relationship, "target_part", None)
            content_type = str(getattr(target, "content_type", ""))
            partname = str(getattr(target, "partname", ""))
            if not content_type.startswith("image/") or partname in seen_parts:
                continue
            if target is None:
                continue
            seen_parts.add(partname)
            filename = _safe_asset_name(partname, f"image-{len(assets) + 1:03d}.bin")
            assets.append(
                GeneratedDocumentFile(
                    relative_path=f"assets/{filename}",
                    content=bytes(target.blob),
                    mime=content_type,
                )
            )
        content = "\n".join(lines).rstrip() + "\n"
        return DocumentConversionPayload(
            status="succeeded",
            files=(
                GeneratedDocumentFile(
                    relative_path="content.md",
                    content=content.encode("utf-8"),
                    mime="text/markdown",
                ),
                *assets,
            ),
        )


def _csv_safe_value(value: object) -> object:
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return f"'{value}"
    return value


def _sheet_filename(title: str, index: int) -> str:
    safe = re.sub(r"[^\w.\-]+", "-", title, flags=re.UNICODE).strip(".-")
    return f"{index:03d}-{(safe or 'sheet')[:100]}"


class XlsxWorkspaceConverter:
    source_extension = ".xlsx"
    converter_name = "workspace-openpyxl"

    @property
    def converter_version(self) -> str:
        return _package_version("openpyxl")

    def convert(self, source_path: Path, *, source_name: str) -> DocumentConversionPayload:
        _zip_preflight(source_path)
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]
        except ImportError as exc:  # pragma: no cover - dependency gate
            raise WorkspaceDocumentDependencyError(
                "openpyxl is required for Workspace XLSX conversion"
            ) from exc
        try:
            workbook = load_workbook(
                source_path,
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
            raise InvalidWorkspaceDocumentError("XLSX could not be parsed") from exc
        if len(workbook.worksheets) > _MAX_WORKBOOK_SHEETS:
            workbook.close()
            raise InvalidWorkspaceDocumentError("workbook has too many sheets")

        outputs: list[GeneratedDocumentFile] = []
        summary = [f"# {_markdown_title(source_name)}", "", "## Workbook summary", ""]
        summary.extend((f"- Sheets: {len(workbook.worksheets)}", ""))
        total_cells = 0
        try:
            for index, sheet in enumerate(workbook.worksheets, start=1):
                buffer = io.StringIO(newline="")
                writer = csv.writer(buffer, lineterminator="\n")
                column_types: dict[int, set[str]] = {}
                row_count = 0
                max_columns = 0
                formula_count = 0
                for row in sheet.iter_rows(values_only=True):
                    row_count += 1
                    max_columns = max(max_columns, len(row))
                    total_cells += len(row)
                    if total_cells > _MAX_WORKBOOK_CELLS:
                        raise InvalidWorkspaceDocumentError(
                            "workbook exceeds the cell processing limit"
                        )
                    safe_row: list[object] = []
                    for column, value in enumerate(row, start=1):
                        if isinstance(value, str) and value.startswith("="):
                            formula_count += 1
                            value_type = "formula"
                        elif value is None:
                            value_type = "empty"
                        else:
                            value_type = type(value).__name__
                        column_types.setdefault(column, set()).add(value_type)
                        safe_row.append(_csv_safe_value(value))
                    writer.writerow(safe_row)

                base = _sheet_filename(sheet.title, index)
                csv_path = f"tables/{base}.csv"
                schema_path = f"tables/{base}.schema.json"
                schema = {
                    "sheet": sheet.title,
                    "range": sheet.calculate_dimension(),
                    "rows": row_count,
                    "columns": max_columns,
                    "formula_cells": formula_count,
                    "formula_policy": "stored_as_inert_text",
                    "column_types": {
                        str(column): sorted(types) for column, types in sorted(column_types.items())
                    },
                }
                outputs.extend(
                    (
                        GeneratedDocumentFile(
                            relative_path=csv_path,
                            content=buffer.getvalue().encode("utf-8"),
                            mime="text/csv",
                        ),
                        GeneratedDocumentFile(
                            relative_path=schema_path,
                            content=(
                                json.dumps(schema, ensure_ascii=False, indent=2, sort_keys=True)
                                + "\n"
                            ).encode("utf-8"),
                            mime="application/json",
                        ),
                    )
                )
                summary.append(
                    f"- **{sheet.title}**: {row_count} rows × {max_columns} columns; "
                    f"{formula_count} formulas; [{base}.csv](tables/{base}.csv)"
                )
        finally:
            workbook.close()
        summary.append("")
        outputs.insert(
            0,
            GeneratedDocumentFile(
                relative_path="content.md",
                content="\n".join(summary).encode("utf-8"),
                mime="text/markdown",
            ),
        )
        return DocumentConversionPayload(status="succeeded", files=tuple(outputs))


@dataclass(frozen=True)
class WorkspaceDocumentConverterRegistry:
    converters: Mapping[str, WorkspaceDocumentConverter]

    @classmethod
    def defaults(cls) -> WorkspaceDocumentConverterRegistry:
        instances: tuple[WorkspaceDocumentConverter, ...] = (
            PypdfWorkspaceConverter(),
            DocxWorkspaceConverter(),
            XlsxWorkspaceConverter(),
        )
        return cls({converter.source_extension: converter for converter in instances})

    def for_extension(self, extension: str) -> WorkspaceDocumentConverter:
        converter = self.converters.get(extension.casefold())
        if converter is None:
            raise UnsupportedWorkspaceDocumentError(
                f"unsupported Workspace document extension {extension!r}"
            )
        return converter


class WorkspaceDocumentService:
    """Snapshot, convert, manifest, and atomically publish a Workspace document."""

    def __init__(
        self,
        store: WorkspaceStore,
        *,
        staging_root: Path,
        registry: WorkspaceDocumentConverterRegistry | None = None,
    ) -> None:
        self._store = store
        self._staging_root = staging_root.resolve(strict=False)
        self._registry = registry or WorkspaceDocumentConverterRegistry.defaults()

    async def convert(
        self, source_file_id: str, session_id: str
    ) -> WorkspaceDocumentConversionResult:
        source = await self._store.get_for_session(session_id, source_file_id)
        if source.purpose != "document_original":
            raise UnsupportedWorkspaceDocumentError(
                "only immutable Workspace document originals can be converted"
            )
        source_path = PurePosixPath(source.logical_path)
        document_root = workspace_document_output_root(source)
        document_id = PurePosixPath(document_root).name
        extension = source_path.suffix.casefold()
        converter = self._registry.for_extension(extension)

        current_refs = await self._store.list_session(session_id)
        existing_manifest = next(
            (
                ref
                for ref in current_refs
                if ref.logical_path.casefold() == f"{document_root}/manifest.json".casefold()
                and ref.purpose == "document_conversion"
            ),
            None,
        )
        self._staging_root.mkdir(parents=True, exist_ok=True)
        conversion_root = self._staging_root / f"conversion-{uuid.uuid4().hex}"
        materialized_root = conversion_root / "workspace"
        try:
            state = await self._store.get_workspace_state(session_id)
            snapshot = await self._store.materialize_workspace_revision(
                session_id,
                materialized_root,
                expected_workspace_revision=state.revision,
            )
            cached = self._read_manifest(existing_manifest)
            if (
                cached is not None
                and self._manifest_matches(cached, source, converter)
                and self._manifest_artifacts_match(cached, current_refs)
            ):
                return await self._result_from_workspace(
                    source,
                    cached,
                    reused=True,
                    workspace_revision=snapshot.revision,
                )
            frozen_source = materialized_root.joinpath(*source_path.parts)
            try:
                payload = await asyncio.to_thread(
                    converter.convert,
                    frozen_source,
                    source_name=source.name,
                )
                manifest = self._success_manifest(
                    source,
                    document_id,
                    converter,
                    payload,
                    document_root,
                )
            except WorkspaceDocumentDependencyError:
                manifest = self._failure_manifest(
                    source,
                    document_id,
                    converter,
                    error_code="dependency_unavailable",
                )
                payload = DocumentConversionPayload(status="succeeded", files=())
            except InvalidWorkspaceDocumentError:
                manifest = self._failure_manifest(
                    source,
                    document_id,
                    converter,
                    error_code="invalid_document",
                )
                payload = DocumentConversionPayload(status="succeeded", files=())
            except WorkspaceDocumentError:
                manifest = self._failure_manifest(
                    source,
                    document_id,
                    converter,
                    error_code="conversion_failed",
                )
                payload = DocumentConversionPayload(status="succeeded", files=())
            except Exception:
                manifest = self._failure_manifest(
                    source,
                    document_id,
                    converter,
                    error_code="conversion_failed",
                )
                payload = DocumentConversionPayload(status="succeeded", files=())

            publish_files: tuple[GeneratedDocumentFile, ...]
            if manifest.status == "failed":
                # Preserve an earlier good conversion. On first failure, only a
                # safe manifest is published; no partial content becomes visible.
                if existing_manifest is not None:
                    return WorkspaceDocumentConversionResult(
                        source_file_id=source.id,
                        document_id=document_id,
                        status="failed",
                        workspace_revision=snapshot.revision,
                        manifest_file_id=existing_manifest.id,
                        warnings=("previous_conversion_preserved",),
                        error_code=manifest.error_code,
                    )
                publish_files = ()
            else:
                publish_files = payload.files
            staged_changes = self._stage_changes(
                conversion_root / "generated",
                document_root,
                publish_files,
                manifest,
            )
            keep_paths = {change.logical_path.casefold() for change in staged_changes}
            deleted_paths = tuple(
                sorted(
                    ref.logical_path
                    for ref in current_refs
                    if ref.purpose == "document_conversion"
                    and ref.logical_path.casefold().startswith(f"{document_root.casefold()}/")
                    and ref.logical_path.casefold() not in keep_paths
                )
            )
            published = await self._store.publish_document_conversion(
                session_id,
                source_file_id=source.id,
                transaction_id=f"publish-{uuid.uuid4().hex}",
                expected_workspace_revision=snapshot.revision,
                expected_workspace_sha256=snapshot.tree_sha256,
                changes=staged_changes,
                deleted_paths=deleted_paths,
            )
            return await self._result_from_workspace(
                source,
                manifest,
                reused=False,
                workspace_revision=published.revision,
            )
        finally:
            shutil.rmtree(conversion_root, ignore_errors=True)

    def _read_manifest(self, ref: FileRef | None) -> WorkspaceDocumentManifest | None:
        if ref is None:
            return None
        try:
            raw = Path(ref.path).read_bytes()
            if len(raw) != ref.size or hashlib.sha256(raw).hexdigest() != ref.sha256:
                return None
            return WorkspaceDocumentManifest.model_validate_json(raw)
        except (OSError, ValueError):
            return None

    @staticmethod
    def _manifest_matches(
        manifest: WorkspaceDocumentManifest,
        source: FileRef,
        converter: WorkspaceDocumentConverter,
    ) -> bool:
        return (
            manifest.source_file_id == source.id
            and manifest.source_sha256 == source.sha256
            and manifest.converter_name == converter.converter_name
            and manifest.converter_version == converter.converter_version
            and manifest.config_version == WORKSPACE_DOCUMENT_CONFIG_VERSION
            and manifest.status in {"succeeded", "needs_ocr"}
        )

    @staticmethod
    def _manifest_artifacts_match(
        manifest: WorkspaceDocumentManifest,
        refs: list[FileRef],
    ) -> bool:
        expected = {
            artifact.logical_path.casefold(): artifact for artifact in manifest.generated_files
        }
        actual = {
            ref.logical_path.casefold(): ref
            for ref in refs
            if ref.purpose == "document_conversion"
            and ref.logical_path.casefold()
            != (f"documents/{manifest.document_id}/manifest.json".casefold())
            and ref.logical_path.casefold().startswith(
                f"documents/{manifest.document_id}/".casefold()
            )
        }
        return expected.keys() == actual.keys() and all(
            ref.size == expected[path].size and ref.sha256 == expected[path].sha256
            for path, ref in actual.items()
        )

    @staticmethod
    def _success_manifest(
        source: FileRef,
        document_id: str,
        converter: WorkspaceDocumentConverter,
        payload: DocumentConversionPayload,
        document_root: str,
    ) -> WorkspaceDocumentManifest:
        artifacts = tuple(
            WorkspaceDocumentArtifact(
                logical_path=f"{document_root}/{item.relative_path}",
                mime=item.mime,
                size=len(item.content),
                sha256=hashlib.sha256(item.content).hexdigest(),
            )
            for item in payload.files
        )
        return WorkspaceDocumentManifest(
            document_id=document_id,
            source_file_id=source.id,
            source_name=source.name,
            source_logical_path=source.logical_path,
            source_mime=source.mime,
            source_size=source.size,
            source_sha256=source.sha256,
            converter_name=converter.converter_name,
            converter_version=converter.converter_version,
            status=payload.status,
            converted_at_ms=int(time.time() * 1000),
            warnings=payload.warnings,
            generated_files=artifacts,
        )

    @staticmethod
    def _failure_manifest(
        source: FileRef,
        document_id: str,
        converter: WorkspaceDocumentConverter,
        *,
        error_code: str,
    ) -> WorkspaceDocumentManifest:
        return WorkspaceDocumentManifest(
            document_id=document_id,
            source_file_id=source.id,
            source_name=source.name,
            source_logical_path=source.logical_path,
            source_mime=source.mime,
            source_size=source.size,
            source_sha256=source.sha256,
            converter_name=converter.converter_name,
            converter_version=converter.converter_version,
            status="failed",
            converted_at_ms=int(time.time() * 1000),
            error_code=error_code,
        )

    @staticmethod
    def _stage_changes(
        staging_dir: Path,
        document_root: str,
        files: tuple[GeneratedDocumentFile, ...],
        manifest: WorkspaceDocumentManifest,
    ) -> tuple[WorkspacePublishChange, ...]:
        all_files = (
            *files,
            GeneratedDocumentFile(
                relative_path="manifest.json",
                content=(manifest.model_dump_json(indent=2) + "\n").encode("utf-8"),
                mime="application/json",
            ),
        )
        if len(all_files) > _MAX_GENERATED_FILES:
            raise InvalidWorkspaceDocumentError("document generated too many files")
        if sum(len(item.content) for item in all_files) > _MAX_GENERATED_BYTES:
            raise InvalidWorkspaceDocumentError("document generated too many bytes")
        seen: set[str] = set()
        changes: list[WorkspacePublishChange] = []
        for index, item in enumerate(all_files):
            logical_path = normalize_workspace_logical_path(f"{document_root}/{item.relative_path}")
            if (
                not is_document_conversion_workspace_path(logical_path, document_root)
                or logical_path.casefold() in seen
            ):
                raise WorkspacePublishPolicyError("converter generated an unsafe output path")
            seen.add(logical_path.casefold())
            staged = staging_dir / f"{index:04d}.blob"
            staged.parent.mkdir(parents=True, exist_ok=True)
            staged.write_bytes(item.content)
            changes.append(
                WorkspacePublishChange(
                    logical_path=logical_path,
                    source_path=staged.resolve(strict=True),
                    size=len(item.content),
                    sha256=hashlib.sha256(item.content).hexdigest(),
                )
            )
        return tuple(changes)

    async def _result_from_workspace(
        self,
        source: FileRef,
        manifest: WorkspaceDocumentManifest,
        *,
        reused: bool,
        workspace_revision: int | None = None,
    ) -> WorkspaceDocumentConversionResult:
        document_root = workspace_document_output_root(source)
        refs = tuple(
            sorted(
                (
                    ref
                    for ref in await self._store.list_session(source.session_id)
                    if ref.purpose == "document_conversion"
                    and ref.logical_path.casefold().startswith(f"{document_root.casefold()}/")
                ),
                key=lambda ref: ref.logical_path,
            )
        )
        manifest_ref = next(
            (ref for ref in refs if ref.logical_path.casefold().endswith("/manifest.json")),
            None,
        )
        primary_ref = next(
            (ref for ref in refs if ref.logical_path.casefold().endswith("/content.md")),
            None,
        )
        if workspace_revision is None:
            workspace_revision = (await self._store.get_workspace_state(source.session_id)).revision
        return WorkspaceDocumentConversionResult(
            source_file_id=source.id,
            document_id=manifest.document_id,
            status=manifest.status,
            reused=reused,
            workspace_revision=workspace_revision,
            manifest_file_id=manifest_ref.id if manifest_ref is not None else None,
            primary_file_id=primary_ref.id if primary_ref is not None else None,
            files=refs,
            warnings=manifest.warnings,
            error_code=manifest.error_code,
        )


__all__ = [
    "WORKSPACE_DOCUMENT_MANIFEST_SCHEMA",
    "WORKSPACE_DOCUMENT_CONFIG_VERSION",
    "DocumentConversionStatus",
    "WorkspaceDocumentError",
    "UnsupportedWorkspaceDocumentError",
    "WorkspaceDocumentDependencyError",
    "InvalidWorkspaceDocumentError",
    "GeneratedDocumentFile",
    "DocumentConversionPayload",
    "WorkspaceDocumentArtifact",
    "WorkspaceDocumentManifest",
    "WorkspaceDocumentConversionResult",
    "WorkspaceDocumentConverter",
    "PypdfWorkspaceConverter",
    "DocxWorkspaceConverter",
    "XlsxWorkspaceConverter",
    "WorkspaceDocumentConverterRegistry",
    "WorkspaceDocumentService",
]
