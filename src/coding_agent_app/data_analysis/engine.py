"""Trusted fixed operations; never evaluates input as Python, SQL or expressions."""

from __future__ import annotations

import csv
import math
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from agent_workspace.analysis import (
    MAX_ANALYSIS_CELLS,
    MAX_ANALYSIS_COLUMNS,
    MAX_ANALYSIS_ROWS,
    MAX_RESULT_ROWS,
    AnalysisError,
    AnalysisRequest,
)


def _check_shape(rows: int, columns: int) -> None:
    if (
        rows > MAX_ANALYSIS_ROWS
        or columns > MAX_ANALYSIS_COLUMNS
        or rows * columns > MAX_ANALYSIS_CELLS
    ):
        raise AnalysisError("data_limit_exceeded")


def _headers(values: list[Any]) -> list[str]:
    names = [str(value).strip() if value is not None else "" for value in values]
    if not names or any(not name or len(name) > 200 for name in names):
        raise AnalysisError("invalid_column_headers")
    if len(set(names)) != len(names):
        raise AnalysisError("duplicate_column_headers")
    return names


def _load(path: Path, request: AnalysisRequest) -> tuple[Any, dict[str, Any], list[str]]:
    import pandas as pd

    extension = path.suffix.lower()
    warnings: list[str] = []
    source: dict[str, Any] = {"format": extension.lstrip("."), "sheet": None}
    records: list[list[Any]] = []
    if extension in {".csv", ".tsv"}:
        if request.sheet is not None:
            raise AnalysisError("sheet_only_supported_for_xlsx")
        csv.field_size_limit(65536)
        with path.open(encoding=request.encoding, newline="") as stream:
            reader = csv.reader(stream, delimiter="\t" if extension == ".tsv" else ",")
            for _ in range(request.header_row):
                next(reader, None)
            headers = _headers(next(reader, []))
            records = []
            for row in reader:
                if not row:
                    continue
                if len(row) != len(headers):
                    raise AnalysisError("inconsistent_column_count")
                _check_shape(len(records) + 1, len(headers))
                records.append([None if cell == "" else cell for cell in row])
        frame = pd.DataFrame(records, columns=headers)
        source["encoding"] = request.encoding
    elif extension == ".xlsx":
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(e.file_size for e in entries) > 256 * 1024**2:
                raise AnalysisError("data_limit_exceeded")
        with path.open("rb") as raw, path.open("rb") as cached:
            formulas = load_workbook(raw, read_only=True, data_only=False, keep_links=False)
            values = load_workbook(cached, read_only=True, data_only=True, keep_links=False)
            try:
                source["sheets"] = formulas.sheetnames
                sheet_name = request.sheet or formulas.sheetnames[0]
                if sheet_name not in formulas.sheetnames:
                    raise AnalysisError("sheet_not_found")
                source["sheet"] = sheet_name
                if request.sheet is None and len(formulas.sheetnames) > 1:
                    warnings.append(
                        "First worksheet selected; specify sheet for another worksheet."
                    )
                sheet = formulas[sheet_name]
                _check_shape(sheet.max_row or 0, sheet.max_column or 0)
                formula_rows = sheet.iter_rows(min_row=request.header_row + 1, values_only=True)
                value_rows = values[sheet_name].iter_rows(
                    min_row=request.header_row + 1,
                    values_only=True,
                )
                headers = _headers(list(next(formula_rows, ())))
                next(value_rows, None)
                records = []
                missing_formula_columns: set[str] = set()
                formula_count = 0
                for frow, vrow in zip(formula_rows, value_rows, strict=True):
                    parsed_row: list[Any] = []
                    for index, (original, value) in enumerate(zip(frow, vrow, strict=True)):
                        if isinstance(original, str) and original.startswith("="):
                            formula_count += 1
                            if value is None:
                                missing_formula_columns.add(headers[index])
                        parsed_row.append(None if value == "" else value)
                    _check_shape(len(records) + 1, len(headers))
                    records.append(parsed_row)
                if formula_count:
                    warnings.append(
                        "Formula values are saved Excel caches, not recalculated values."
                    )
                source["formula_cells"] = formula_count
                source["missing_formula_columns"] = sorted(missing_formula_columns)
                if missing_formula_columns:
                    warnings.append(
                        "Some formulas have no cached value. Recalculate in Excel first."
                    )
                frame = pd.DataFrame(records, columns=headers)
            finally:
                formulas.close()
                values.close()
    elif extension == ".parquet":
        import pyarrow.parquet as pq

        if request.sheet is not None or request.header_row:
            raise AnalysisError("invalid_parquet_options")
        parquet = pq.ParquetFile(path)
        try:
            _check_shape(parquet.metadata.num_rows, len(parquet.schema_arrow))
            expanded = sum(
                parquet.metadata.row_group(i).total_byte_size
                for i in range(parquet.metadata.num_row_groups)
            )
            if expanded > 256 * 1024**2:
                raise AnalysisError("data_limit_exceeded")
            frame = parquet.read().to_pandas()
            frame.columns = _headers(list(frame.columns))
        finally:
            parquet.close()
    else:
        raise AnalysisError("unsupported_format")
    _check_shape(len(frame), len(frame.columns))
    if any(
        any(isinstance(v, (dict, list, tuple, bytes)) for v in frame[c].dropna())
        for c in frame.columns
    ):
        raise AnalysisError("nested_columns_not_supported")
    source["header_row"] = request.header_row
    return frame, source, warnings


def _number(series: Any) -> Any:
    import pandas as pd

    try:
        result = pd.to_numeric(series, errors="raise")
        if any(not math.isfinite(float(v)) for v in result.dropna()):
            raise ValueError("nonfinite")
        return result
    except (ValueError, TypeError, OverflowError) as exc:
        raise AnalysisError("invalid_numeric_values") from exc


def _scalar(value: Any) -> Any:
    import pandas as pd

    if pd.isna(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, int) and abs(value) > 2**53 - 1:
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _inferred_type(series: Any) -> str:
    values = series.dropna()
    if not len(values):
        return "empty"
    if all(isinstance(v, bool) for v in values):
        return "boolean"
    # Do not erase identifier leading zeros or infer dates from arbitrary strings.
    if all(
        isinstance(v, (int, float))
        or (isinstance(v, str) and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", v))
        for v in values
    ):
        return "number"
    if all(isinstance(v, (datetime, date)) for v in values):
        return "datetime"
    return "text"


def _render_chart(frame: Any, request: AnalysisRequest, output: Path) -> None:
    import matplotlib

    matplotlib.use("Agg")
    matplotlib.rcParams.update({"text.usetex": False, "text.parse_math": False})
    from matplotlib.figure import Figure

    figure = Figure(figsize=(8, 4), layout="constrained")
    axes = figure.subplots()
    x, y = str(request.x), str(request.y)
    if request.chart_type == "histogram":
        axes.bar(frame["bin"].astype(str), frame["count"])
        axes.tick_params(axis="x", labelrotation=45)
        axes.set_ylabel("count")
    elif request.chart_type == "scatter":
        axes.scatter(frame[x], frame[y], s=14)
        axes.set_ylabel(y)
    else:
        labels = [str(v)[:60] for v in frame[x]]
        if request.chart_type == "bar":
            axes.bar(labels, frame[y])
        else:
            axes.plot(labels, frame[y], marker=".")
        axes.tick_params(axis="x", labelrotation=45)
        axes.set_ylabel(y[:100])
    axes.set_xlabel(x[:100])
    figure.savefig(output, format="png", dpi=120)
    figure.clear()


def calculate(path: Path, request: AnalysisRequest, output_dir: Path) -> dict[str, Any]:
    import pandas as pd

    frame, source, warnings = _load(path, request)
    source_rows = len(frame)
    schema = [{"name": c, "type": _inferred_type(frame[c])} for c in frame.columns]
    referenced = [
        *request.columns,
        *request.group_by,
        *(f.column for f in request.filters),
        *(m.column for m in request.metrics),
    ]
    referenced.extend(c for c in [request.date_column, request.x, request.y] if c is not None)
    if any(c not in frame.columns for c in referenced):
        raise AnalysisError("column_not_found")
    missing_formulas = source.get("missing_formula_columns", [])
    used_columns = referenced or list(frame.columns)
    if request.action == "profile" and not request.columns:
        used_columns = list(frame.columns)
    if request.action != "inspect" and set(used_columns).intersection(missing_formulas):
        raise AnalysisError("formula_cache_missing")
    for condition in request.filters:
        series = frame[condition.column]
        operation, value = condition.operation, condition.value
        if operation in {"gt", "gte", "lt", "lte"} or isinstance(value, (int, float)):
            series = _number(series)
        if operation == "is_null":
            mask = series.isna()
        elif operation == "not_null":
            mask = series.notna()
        else:
            method = {"eq": "eq", "ne": "ne", "gt": "gt", "gte": "ge", "lt": "lt", "lte": "le"}[
                operation
            ]
            try:
                mask = getattr(series, method)(value).fillna(False) & series.notna()
            except (ValueError, TypeError) as exc:
                raise AnalysisError("invalid_filter_value") from exc
        frame = frame.loc[mask].copy()
    analyzed_rows = len(frame)
    duplicate_rows = int(frame.duplicated().sum())
    if request.action == "inspect":
        result = frame[request.columns or list(frame.columns)]
    elif request.action == "profile":
        records = []
        for column in request.columns or list(frame.columns):
            series = frame[column]
            kind = _inferred_type(series)
            stats: dict[str, Any] = {
                "column": column,
                "type": kind,
                "non_null": int(series.count()),
                "missing": int(series.isna().sum()),
                "unique": int(series.nunique()),
            }
            if kind == "number":
                numbers = _number(series)
                stats.update(
                    {
                        "min": numbers.min(),
                        "max": numbers.max(),
                        "mean": numbers.mean(),
                        "p25": numbers.quantile(0.25),
                        "median": numbers.median(),
                        "p75": numbers.quantile(0.75),
                    }
                )
            records.append(stats)
        result = pd.DataFrame(records)
    elif request.action in {"aggregate", "timeseries"}:
        groups = list(request.group_by)
        if request.action == "timeseries":
            try:
                parsed = pd.to_datetime(
                    frame[request.date_column], format=request.date_format, errors="raise", utc=True
                )
                if parsed.isna().any():
                    raise ValueError("null date")
                frequency = {"day": "D", "week": "W-SUN", "month": "M", "year": "Y"}
                if "period" in frame.columns:
                    raise AnalysisError("reserved_period_column")
                frame["period"] = (
                    parsed.dt.tz_localize(None)
                    .dt.to_period(
                        frequency[request.frequency],
                    )
                    .astype(str)
                )
                groups = ["period", *groups]
                warnings.append("Time buckets use UTC; dates are parsed with the requested format.")
            except (ValueError, TypeError, AttributeError) as exc:
                raise AnalysisError("invalid_datetime_values") from exc
        named: dict[str, Any] = {}
        for index, metric in enumerate(request.metrics):
            alias = f"{metric.column}__{metric.operation}"
            if alias in groups:
                raise AnalysisError("result_column_conflict")
            column = metric.column
            if metric.operation != "count":
                column = f"__analysis_metric_{index}"
                while column in frame.columns:
                    column += "_"
                frame[column] = _number(frame[metric.column])
            function = (
                (lambda s: s.sum(min_count=1)) if metric.operation == "sum" else metric.operation
            )
            named[alias] = pd.NamedAgg(column=column, aggfunc=function)
        if groups:
            result = frame.groupby(groups, dropna=False, sort=True).agg(**named).reset_index()
        else:
            result = pd.DataFrame(
                [
                    {
                        alias: function(frame[column])
                        if callable(function)
                        else getattr(frame[column], function)()
                        for alias, (column, function) in named.items()
                    }
                ]
            )
    else:
        x, y = str(request.x), str(request.y)
        if request.chart_type == "histogram":
            import numpy as np

            numbers = _number(frame[x]).dropna()
            counts, bins = np.histogram(numbers, bins=10)
            result = pd.DataFrame(
                {
                    "bin": [f"{a:.4g}–{b:.4g}" for a, b in zip(bins[:-1], bins[1:], strict=True)],
                    "count": counts,
                }
            )
        else:
            frame[y] = _number(frame[y])
            if request.chart_type == "scatter":
                frame[x] = _number(frame[x])
                result = frame[[x, y]].dropna()
            else:
                result = frame.groupby(x, dropna=False, sort=True)[y].sum(min_count=1).reset_index()
                warnings.append(
                    "Chart groups equal X values and sums Y; null Y values are excluded."
                )
    if request.sort_by is not None:
        if request.sort_by not in result.columns:
            raise AnalysisError("sort_column_not_found")
        result = result.sort_values(
            request.sort_by,
            ascending=not request.descending,
            kind="stable",
            key=_number if _inferred_type(result[request.sort_by]) == "number" else None,
        )
    total_result_rows = len(result)
    limit = min(request.limit, 50) if request.action == "chart" else request.limit
    result = result.head(min(limit, MAX_RESULT_ROWS))
    if len(result) < total_result_rows:
        warnings.append(
            f"Result limited to {len(result)} of {total_result_rows} rows; not a full export."
        )
    if request.action == "chart":
        _render_chart(result, request, output_dir / "chart.png")
    rows = [[_scalar(v) for v in row] for row in result.itertuples(index=False, name=None)]
    with (output_dir / "result.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)

        def safe(value: Any) -> Any:
            return (
                "'" + value
                if isinstance(value, str)
                and value.lstrip()[:1]
                in {
                    "=",
                    "+",
                    "-",
                    "@",
                }
                else value
            )

        writer.writerow([safe(str(c)) for c in result.columns])
        writer.writerows([safe(v) for v in row] for row in rows)
    return {
        "source": source,
        "source_rows": source_rows,
        "analyzed_rows": analyzed_rows,
        "schema": schema,
        "columns": list(result.columns),
        "rows": rows[:50],
        "result_rows": total_result_rows,
        "exported_rows": len(rows),
        "preview_rows": min(50, len(rows)),
        "limited": len(rows) < total_result_rows,
        "preview_limited": len(rows) > 50,
        "duplicate_rows": duplicate_rows,
        "warnings": warnings,
        "null_policy": "empty cells are null; aggregates exclude nulls; count counts non-nulls",
        "numeric_policy": "pandas numeric types; not exact-decimal accounting",
        "has_chart": request.action == "chart",
    }
